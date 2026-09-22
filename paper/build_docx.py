"""Build paper/paper.docx from paper/paper.md (the source of truth).

Handles only the Markdown subset used in paper.md: #/##/### headings, paragraphs with
**bold** / *italic* / `code`, bullet and numbered lists, pipe tables, images, and the
leading <!-- draft notes --> comment (shown as a grey note). Image syntax:

    ![caption](figures/a.png)                 one figure, caption below
    ![caption](figures/a.png;figures/b.png)   several images side by side in one row
    ![caption](figures/a.png){4.2}            optional maximum width in inches

Uses python-docx, which is installed for the system Python (not the project venv):

    python3 paper/build_docx.py                          # paper.md -> paper.docx (default)
    python3 paper/build_docx.py SRC.md OUT.docx           # any other Markdown file, same rules
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
SRC, OUT = HERE / "paper.md", HERE / "paper.docx"
FONT, BODY_PT = "Times New Roman", 10.5
TEXT_WIDTH = 8.5 - 2 * 0.9            # inches between the margins
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")
IMAGE = re.compile(r"!\[(.*?)\]\((.*?)\)(?:\{(\d+(?:\.\d+)?)\})?")


def set_font(run, size=None, bold=None, italic=None, colour=None, name=FONT) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if colour is not None:
        run.font.color.rgb = colour


def add_runs(paragraph, text: str, size: float = None, colour: RGBColor = None, bold_all: bool = False) -> None:
    """Add `text` to `paragraph`, turning **bold**, *italic* and `code` into formatted runs."""
    text = text.replace("&nbsp;", " ")
    for part in INLINE.split(text):
        if not part:
            continue
        bold = italic = False
        name = FONT
        if part.startswith("**") and part.endswith("**"):
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, name = part[1:-1], "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        run = paragraph.add_run(part)
        set_font(run, size=size, bold=bold or bold_all, italic=italic, colour=colour, name=name)


def tight(paragraph, before=0, after=4, keep_next=False, align=None) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), 1.0
    fmt.keep_with_next = keep_next
    if align is not None:
        paragraph.alignment = align


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_table(doc: Document, rows: list) -> None:
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate([header] + body):
        for c, cell_text in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            para = cell.paragraphs[0]
            tight(para, after=0, keep_next=r < len(body))
            add_runs(para, cell_text, size=8.5, bold_all=(r == 0))
            if r == 0:
                shade(cell, "E7E6E6")
    tight(doc.add_paragraph(), after=2)


def add_figure(doc: Document, caption: str, paths: list, max_width: float = None) -> None:
    n = len(paths)
    holder = doc.add_table(rows=1, cols=n) if n > 1 else None
    col = min(max_width or TEXT_WIDTH, TEXT_WIDTH) / n - (0.08 if n > 1 else 0)
    for k, rel in enumerate(paths):
        path = (HERE / rel.strip()).resolve()
        with Image.open(path) as im:
            aspect = im.width / im.height
        max_h = 7.6 if aspect < 0.6 else 4.2
        width = min(col, max_h * aspect)
        para = holder.rows[0].cells[k].paragraphs[0] if holder is not None else doc.add_paragraph()
        tight(para, before=2, after=2, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        para.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        cap = doc.add_paragraph()
        tight(cap, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_runs(cap, f"*{caption}*", size=9)


def add_page_number_footer(doc: Document) -> None:
    para = doc.sections[0].footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    set_font(run, size=9)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


def main() -> None:
    global HERE
    src = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else SRC
    out = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else OUT
    HERE = src.parent   # image paths in the Markdown are resolved relative to its own folder
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin, section.bottom_margin = Inches(0.85), Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for level, size in ((1, 13), (2, 11.5), (3, 10.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = FONT, Pt(size), True
        style.font.color.rgb = RGBColor(0, 0, 0)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            style.element.rPr.rFonts.set(qn(attr), FONT)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(9 if level == 1 else 6), Pt(3)
        style.paragraph_format.keep_with_next = True
    add_page_number_footer(doc)

    text = src.read_text(encoding="utf-8")
    grey = RGBColor(0x66, 0x66, 0x66)
    note = re.match(r"<!--(.*?)-->", text, re.S)
    if note:
        p = doc.add_paragraph()
        tight(p, after=6)
        add_runs(p, "*DRAFT NOTES (delete before submission):* " + " ".join(note.group(1).split()), size=8, colour=grey)
        text = text[note.end():]

    lines, i, first_h1, in_references = text.splitlines(), 0, True, False
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped or stripped == "---":
            i += 1
        elif stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, [r for r in block if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)])
        elif re.match(r"^#{1,3} ", stripped):
            level = len(stripped.split()[0])
            if level == 1 and first_h1:          # document title
                p = doc.add_paragraph()
                tight(p, before=4, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
                add_runs(p, stripped[2:], size=16, bold_all=True)
                first_h1 = False
            else:
                add_runs(doc.add_heading(level=level - 1), stripped[level + 1:], size={1: 13, 2: 11.5, 3: 10.5}[level], colour=RGBColor(0, 0, 0), bold_all=True)
                in_references = stripped[level + 1:].strip() == "References"
            i += 1
        elif stripped.startswith("!["):
            m = IMAGE.match(stripped)
            add_figure(doc, m.group(1), m.group(2).split(";"), float(m.group(3)) if m.group(3) else None)
            i += 1
        elif re.match(r"^(- |\d+\. )", stripped):
            # Ordered items are numbered by hand (literal "N. " text), not Word's built-in numbered-list
            # style: that style shares one running counter document-wide, so a second numbered list (e.g.
            # a later "1., 2., 3." block) would silently continue from the first list's last number.
            ordered = re.match(r"^(\d+)\. ", stripped)
            body = stripped[ordered.end():] if ordered else stripped[2:]
            p = doc.add_paragraph()
            tight(p, after=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            add_runs(p, f"{ordered.group(1)}. {body}" if ordered else f"• {body}")
            i += 1
        else:
            para = [stripped]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3} |\||!\[|- |\d+\. |---)", lines[i].strip()):
                para.append(lines[i].strip())
                i += 1
            joined = " ".join(para)
            p = doc.add_paragraph()
            if joined.startswith("**Table "):     # table caption: sits above, stays with its table
                tight(p, before=4, after=2, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)
                add_runs(p, joined, size=9)
            elif joined.startswith("**Author:**"):
                tight(p, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
                add_runs(p, joined, size=10)
            elif in_references and not joined.startswith("*"):   # APA reference entry: hanging indent
                tight(p, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
                add_runs(p, joined)
            else:
                tight(p, after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                add_runs(p, joined)

    doc.save(out)
    print(f"Wrote {out} ({out.stat().st_size // 1024} KiB)")


if __name__ == "__main__":
    main()
